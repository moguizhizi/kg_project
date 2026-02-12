from neo4j import GraphDatabase
from openpyxl import Workbook
from itertools import islice
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

URI = "bolt://localhost:7687"
USER = "neo4j"
PASSWORD = "password"

OUTPUT_PATH = "/home/temp/dataset/home_based_user_training_20260123_v2/analysis/game_task_stats_with_time.xlsx"


def batch_iterable(iterable, batch_size):
    it = iter(iterable)
    while True:
        batch = list(islice(it, batch_size))
        if not batch:
            break
        yield batch


#
QUERY = """
MATCH (g:Game)
WHERE g.id IN $game_ids

MATCH (g)<-[:游戏]-(t:TaskInstance)<-[:包含]-(set:TaskInstanceSet)

WITH
    g,
    toFloatOrNull(set.执行年龄) AS age,
    t.结果 AS result,
    date(set.训练日期) AS train_date

WHERE age IS NOT NULL
AND train_date IS NOT NULL

WITH
    g,
    CASE
        WHEN age < $child_age THEN "Child"
        ELSE "NonChild"
    END AS age_group,
    result,
    train_date

WITH
    g,
    age_group,
    COUNT(*) AS total,
    SUM(CASE WHEN result = "完成" THEN 1 ELSE 0 END) AS completed,
    MIN(train_date) AS start_date,
    MAX(train_date) AS end_date

RETURN
    age_group,
    g.id AS game_id,
    g.name AS game_name,
    total,
    completed,
    ROUND(100.0 * completed / total, 2) AS completed_pct,
    start_date,
    end_date
"""


def fetch_game_ids(tx):
    result = tx.run("MATCH (g:Game) RETURN g.id AS id ORDER BY id")
    return [r["id"] for r in result]


def append_row(sheet, row):
    sheet.append(
        [
            row["game_id"],
            row["game_name"],
            row["total"],
            row["completed"],
            row["completed_pct"],
            str(row["start_date"]),
            str(row["end_date"]),
        ]
    )


def main():

    driver = GraphDatabase.driver(URI, auth=(USER, PASSWORD))

    wb = Workbook()
    wb.remove(wb.active)

    sheets = {
        "Child": wb.create_sheet("Child"),
        "NonChild": wb.create_sheet("NonChild"),
    }

    headers = [
        "game_id",
        "game_name",
        "total",
        "completed",
        "completed_pct",
        "start_date",
        "end_date",
    ]

    for s in sheets.values():
        s.append(headers)

    try:
        with driver.session() as session:

            logger.info("Fetching all game ids...")
            game_ids = session.execute_read(fetch_game_ids)

            logger.info(f"Total games: {len(game_ids)}")

            batch_size = 50  #

            for i, batch in enumerate(batch_iterable(game_ids, batch_size)):

                result = session.run(QUERY, game_ids=batch, child_age=12)

                count = 0

                for record in result:
                    age_group = record["age_group"]

                    if age_group in sheets:
                        append_row(sheets[age_group], record)
                        count += 1

                logger.info(f"Batch {i+1} | games={len(batch)} | rows={count}")

        wb.save(OUTPUT_PATH)
        logger.info(f"Excel saved to: {OUTPUT_PATH}")

    finally:
        driver.close()


if __name__ == "__main__":
    main()
