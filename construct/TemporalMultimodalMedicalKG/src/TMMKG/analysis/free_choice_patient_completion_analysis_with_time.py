from TMMKG.graph.neo4j_db import fetch_node_ids
from TMMKG.infra.neo4j_db import create_neo4j_driver
from itertools import islice
from openpyxl import Workbook
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

URI = "bolt://localhost:7687"
USER = "neo4j"
PASSWORD = "password"

driver = create_neo4j_driver(uri=URI, user=USER, password=PASSWORD)

OUTPUT_XLSX = "/home/temp/dataset/home_based_user_training_20260123_v2/analysis/patient_free_choice_stats_with_time.xlsx"


QUERY = """
MATCH (p:Patient)
WITH p
ORDER BY p.id
SKIP $skip
LIMIT $limit

MATCH (p)-[:参加]->(e:TaskInstanceSet)-[:包含]->(t:TaskInstance)
WHERE toFloatOrNull(e.执行年龄) IS NOT NULL
AND e.训练日期 IS NOT NULL

WITH
    p.id AS patient_id,
    e.id AS event_id,
    toFloat(e.执行年龄) AS age,
    date(e.训练日期) AS train_date,
    MAX(t.任务类型 = "自由") AS has_free,
    MAX(t.任务类型 = "专属") AS has_exclusive

WITH
    patient_id,
    event_id,
    train_date,
    has_free,
    has_exclusive,
    CASE
        WHEN age < $child_age THEN "Child"
        ELSE "NonChild"
    END AS age_group

WITH
    patient_id,
    age_group,
    SUM(CASE WHEN has_free THEN 1 ELSE 0 END) AS free_events,
    SUM(CASE WHEN has_exclusive THEN 1 ELSE 0 END) AS exclusive_events,
    COUNT(*) AS total_events,
    MIN(CASE WHEN has_free THEN train_date END) AS free_start_date,
    MAX(CASE WHEN has_free THEN train_date END) AS free_end_date

RETURN
    age_group,
    patient_id,
    total_events,
    free_events,
    exclusive_events,
    CASE 
        WHEN exclusive_events = 0 THEN NULL
        ELSE ROUND(1.0 * free_events / exclusive_events, 2)
    END AS free_to_exclusive_ratio,
    free_start_date,
    free_end_date

ORDER BY patient_id
"""


def batch_iterable(iterable, batch_size):
    it = iter(iterable)
    while True:
        batch = list(islice(it, batch_size))
        if not batch:
            break
        yield batch


def create_excel():
    wb = Workbook()
    wb.remove(wb.active)

    sheets = {
        "Child": wb.create_sheet("Child"),
        "NonChild": wb.create_sheet("NonChild"),
    }

    headers = [
        "patient_id",
        "total_events",
        "free_events",
        "exclusive_events",
        "free_to_exclusive_ratio",
        "free_start_date",
        "free_end_date",
    ]

    for sheet in sheets.values():
        sheet.append(headers)

    return wb, sheets


def append_row(sheet, row):
    sheet.append(
        [
            row["patient_id"],
            row["total_events"],
            row["free_events"],
            row["exclusive_events"],
            row["free_to_exclusive_ratio"],
            str(row["free_start_date"]) if row["free_start_date"] else None,
            str(row["free_end_date"]) if row["free_end_date"] else None,
        ]
    )


def run_patient_stats(driver):

    wb, sheets = create_excel()

    batch_size = 1000
    params_base = {"child_age": 12}

    try:
        with driver.session() as session:

            all_patient_ids = session.execute_read(fetch_node_ids, "Patient")
            logger.info(f"Total Patient nodes: {len(all_patient_ids)}")

            skip = 0

            for batch in batch_iterable(all_patient_ids, batch_size):

                params = {**params_base, "skip": skip, "limit": len(batch)}

                result = session.run(QUERY, params)

                record_count = 0

                for record in result:
                    age_group = record["age_group"]

                    if age_group not in sheets:
                        continue

                    append_row(sheets[age_group], record)
                    record_count += 1

                logger.info(
                    f"Processed patients {skip} ~ {skip + len(batch)} | records={record_count}"
                )

                skip += len(batch)

    finally:
        wb.save(OUTPUT_XLSX)
        driver.close()

    logger.info(f"Excel saved to: {OUTPUT_XLSX}")


# =====================
# RUN
# =====================

run_patient_stats(driver)
