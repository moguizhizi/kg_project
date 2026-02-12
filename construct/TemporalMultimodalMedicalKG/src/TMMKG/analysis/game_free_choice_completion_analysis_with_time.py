from TMMKG.graph.neo4j_db import fetch_node_ids
from TMMKG.infra.neo4j_db import create_neo4j_driver
from openpyxl import Workbook
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

URI = "bolt://localhost:7687"
USER = "neo4j"
PASSWORD = "password"

OUTPUT_XLSX = "/home/temp/dataset/home_based_user_training_20260123_v2/analysis/game_free_choice_stats_with_time.xlsx"

driver = create_neo4j_driver(uri=URI, user=USER, password=PASSWORD)

###################################
# Cypher
###################################

QUERY = """
MATCH (set:TaskInstanceSet)-[:包含]->(t:TaskInstance)-[:游戏]->(g:Game {id:$game_id})
WHERE 
    toFloatOrNull(set.执行年龄) IS NOT NULL
    AND set.训练日期 IS NOT NULL

WITH 
    g,
    toFloat(set.执行年龄) AS age,
    t.结果 AS result,
    t.任务类型 AS task_type,
    date(set.训练日期) AS train_date

WITH
    g,
    CASE
        WHEN age < $child_age THEN "Child"
        ELSE "NonChild"
    END AS age_group,
    result,
    task_type,
    train_date

WITH
    g,
    age_group,
    COUNT(*) AS total,
    SUM(CASE WHEN task_type = "自由" THEN 1 ELSE 0 END) AS free_total,
    SUM(CASE 
        WHEN task_type = "自由" AND result = "完成"
        THEN 1 ELSE 0 
    END) AS free_completed,

    MIN(CASE 
        WHEN task_type = "自由" AND result = "完成"
        THEN train_date
    END) AS start_date,

    MAX(CASE 
        WHEN task_type = "自由" AND result = "完成"
        THEN train_date
    END) AS end_date

WITH
    g,
    age_group,
    total,
    free_total,
    free_completed,
    start_date,
    end_date,

    CASE 
        WHEN free_total = 0 THEN 0.0
        ELSE ROUND(100.0 * free_completed / free_total, 2)
    END AS free_completed_pct,

    CASE 
        WHEN total = 0 THEN 0.0
        ELSE ROUND(100.0 * free_completed / total, 2)
    END AS free_completed_over_total_pct

RETURN 
    age_group,
    g.id AS game_id,
    g.name AS game_name,
    total,
    free_total,
    free_completed,
    free_completed_pct,
    free_completed_over_total_pct,
    start_date,
    end_date
"""

###################################
# 查询函数
###################################


def get_taskinstances_by_game_result(tx, game_id, child_age=12):
    result = tx.run(
        QUERY,
        game_id=game_id,
        child_age=child_age,
    )
    return [record.data() for record in result]


###################################
# Excel初始化
###################################


def create_workbook():

    wb = Workbook()

    # 删除默认sheet
    wb.remove(wb.active)

    sheets = {
        "Child": wb.create_sheet("Child"),
        "NonChild": wb.create_sheet("NonChild"),
    }

    headers = [
        "game_id",
        "game_name",
        "total",
        "free_total",
        "free_completed",
        "free_completed_over_total_pct",
        "free_completed_pct",
        "free_completed_start_date",
        "free_completed_end_date",
    ]

    for sheet in sheets.values():
        sheet.append(headers)

    return wb, sheets


###################################
# 主逻辑
###################################


def run_game_stats(driver):

    wb, sheets = create_workbook()

    with driver.session() as session:

        all_game_ids = session.execute_read(fetch_node_ids, "Game")
        logger.info(f"Total Game nodes: {len(all_game_ids)}")

        for idx, game_id in enumerate(all_game_ids, 1):

            rows = session.execute_read(
                get_taskinstances_by_game_result,
                game_id,
            )

            for row in rows:

                age_group = row["age_group"]

                if age_group not in sheets:
                    continue

                sheets[age_group].append(
                    [
                        row["game_id"],
                        row["game_name"],
                        row["total"],
                        row["free_total"],
                        row["free_completed"],
                        row["free_completed_over_total_pct"],
                        row["free_completed_pct"],
                        str(row["start_date"]) if row["start_date"] else None,
                        str(row["end_date"]) if row["end_date"] else None,
                    ]
                )

            if idx % 50 == 0:
                logger.info(f"Processed {idx} games")

    wb.save(OUTPUT_XLSX)
    logger.info(f"Excel saved -> {OUTPUT_XLSX}")


###################################
# 启动
###################################

run_game_stats(driver)
driver.close()
