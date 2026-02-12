from TMMKG.graph.neo4j_db import fetch_node_ids
from TMMKG.infra.neo4j_db import create_neo4j_driver
from itertools import islice
from openpyxl import Workbook
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Neo4j 连接信息
URI = "bolt://localhost:7687"
USER = "neo4j"
PASSWORD = "password"

OUTPUT_XLSX = "/home/temp/dataset/home_based_user_training_20260123_v2/analysis/patient_stats_with_time.xlsx"

driver = create_neo4j_driver(uri=URI, user=USER, password=PASSWORD)


def batch_iterable(iterable, batch_size):
    """将可迭代对象分批生成"""
    it = iter(iterable)
    while True:
        batch = list(islice(it, batch_size))
        if not batch:
            break
        yield batch


QUERY = """
MATCH (p:Patient)
WITH p
ORDER BY p.id
SKIP $skip
LIMIT $limit

MATCH (p)-[:参加]->(e:TaskInstanceSet)-[:包含]->(t:TaskInstance)
WHERE 
    toFloatOrNull(e.执行年龄) IS NOT NULL
    AND e.训练日期 IS NOT NULL

WITH 
    p.id AS patient_id,
    e,
    toFloat(e.执行年龄) AS age,
    date(e.训练日期) AS train_date,
    COUNT(t) AS total_tasks,
    SUM(CASE WHEN t.结果 = '完成' THEN 1 ELSE 0 END) AS completed_tasks

WITH
    patient_id,
    age,
    train_date,
    CASE 
        WHEN total_tasks = completed_tasks THEN 1
        ELSE 0
    END AS completed_flag

WITH
    patient_id,
    train_date,
    completed_flag,
    CASE 
        WHEN age < $child_age THEN "Child"
        ELSE "NonChild"
    END AS age_group

WITH
    patient_id,
    age_group,
    COUNT(*) AS total_events,
    SUM(completed_flag) AS completed_events,
    MIN(train_date) AS start_date,
    MAX(train_date) AS end_date

RETURN
    age_group,
    patient_id,
    total_events,
    completed_events,
    ROUND(100.0 * completed_events / total_events, 2) AS completion_pct,
    start_date,
    end_date

ORDER BY patient_id
"""


def create_workbook():
    wb = Workbook()
    wb.remove(wb.active)

    sheets = {
        "Child": wb.create_sheet("Child"),
        "NonChild": wb.create_sheet("NonChild"),
    }

    headers = [
        "patient_id",
        "total_events",
        "completed_events",
        "completion_pct",
        "start_date",
        "end_date",
    ]

    for sheet in sheets.values():
        sheet.append(headers)

    return wb, sheets


def append_row(sheet, record):
    sheet.append(
        [
            record["patient_id"],
            record["total_events"],
            record["completed_events"],
            record["completion_pct"],
            str(record["start_date"]),
            str(record["end_date"]),
        ]
    )


def run_patient_stats(driver):

    wb, sheets = create_workbook()

    batch_size = 1000
    params_base = {"child_age": 12}

    try:
        with driver.session() as session:

            all_patient_ids = session.execute_read(fetch_node_ids, "Patient")
            logger.info(f"Total Patient nodes: {len(all_patient_ids)}")

            skip = 0

            for batch in batch_iterable(all_patient_ids, batch_size):

                params = {**params_base, "skip": skip, "limit": len(batch)}

                result = session.run(QUERY, params)

                count = 0

                for record in result:

                    age_group = record["age_group"]

                    if age_group not in sheets:
                        continue

                    append_row(sheets[age_group], record)
                    count += 1

                logger.info(
                    f"Processed patients {skip} ~ {skip + len(batch)} | rows={count}"
                )

                skip += len(batch)

    finally:
        wb.save(OUTPUT_XLSX)
        logger.info(f"Excel saved to: {OUTPUT_XLSX}")


if __name__ == "__main__":
    run_patient_stats(driver)
